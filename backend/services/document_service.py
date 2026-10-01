from io import BytesIO
from pathlib import Path
from typing import Optional
import re
from html import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

from backend.services.gemini_generator import GeminiDocumentGenerator


_generator = GeminiDocumentGenerator()


def sanitize_text(text: str) -> str:
    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u00a0": " ",
        "\u2022": "-",
        "\ufeff": "",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    return text.strip()


def parse_terms(terms: str) -> list[str]:
    parts = re.split(r";|\n", terms)
    cleaned = []

    for part in parts:
        part = re.sub(r"^\s*[-•*]\s*", "", part).strip()

        if part:
            cleaned.append(part)

    return cleaned


def split_document_lines(text: str) -> list[str]:
    clean_text = sanitize_text(text)

    return [
        line.strip()
        for line in clean_text.splitlines()
        if line.strip()
    ]


def make_txt(text: str) -> bytes:
    return sanitize_text(text).encode("utf-8")


def is_heading(line: str) -> bool:
    if len(line) > 100:
        return False

    heading_names = {
        "title",
        "parties",
        "effective date",
        "purpose",
        "definitions",
        "term",
        "termination",
        "confidentiality",
        "payment",
        "obligations",
        "notices",
        "governing law",
        "dispute resolution",
        "signatures",
        "ai draft notice",
        "intellectual property",
        "representations",
        "warranties",
        "amendments",
        "severability",
        "entire agreement",
    }

    return (
        line.isupper()
        or bool(re.match(r"^\d+[.)]\s+", line))
        or line.lower() in heading_names
    )


def generate_document(
    document_type: str,
    parties: str,
    terms: str,
    effective_date: str,
    brand_name: Optional[str] = None,
) -> str:
    return _generator.generate_document(
        document_type=document_type,
        parties=parties,
        terms=terms,
        effective_date=effective_date,
        brand_name=brand_name,
    )


def make_docx(
    text: str,
    doc_type: str,
    brand_name: Optional[str] = None,
    logo_path: Optional[str] = None,
    terms: Optional[str] = None,
) -> bytes:
    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(11)

    if logo_path and Path(logo_path).exists():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        paragraph.add_run().add_picture(
            logo_path,
            width=Inches(1.4),
        )

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_run = title.add_run(doc_type.upper())
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(16)

    if brand_name:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run(brand_name)
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)

    document.add_paragraph()

    for line in split_document_lines(text):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)

        run = paragraph.add_run(line)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)

        if is_heading(line):
            run.bold = True

    if terms:
        parsed_terms = parse_terms(terms)

        if parsed_terms:
            document.add_paragraph()

            heading = document.add_paragraph()
            heading_run = heading.add_run("KEY USER-PROVIDED TERMS")
            heading_run.bold = True
            heading_run.font.name = "Times New Roman"
            heading_run.font.size = Pt(12)

            table = document.add_table(rows=1, cols=2)
            table.style = "Table Grid"

            table.rows[0].cells[0].text = "No."
            table.rows[0].cells[1].text = "Term"

            for index, term in enumerate(parsed_terms, start=1):
                cells = table.add_row().cells
                cells[0].text = str(index)
                cells[1].text = term

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer_run = footer.add_run(
        f"LegalEase | "
        f"{brand_name or 'AI-assisted draft'} | "
        "Review by a qualified legal professional before use"
    )

    footer_run.font.size = Pt(8)

    buffer = BytesIO()
    document.save(buffer)

    return buffer.getvalue()


def pdf_header_footer(
    canvas,
    document,
    brand_name: Optional[str],
    logo_path: Optional[str],
):
    canvas.saveState()

    width, height = A4

    if logo_path and Path(logo_path).exists():
        try:
            canvas.drawImage(
                logo_path,
                width / 2 - 15 * mm,
                height - 25 * mm,
                width=30 * mm,
                height=15 * mm,
                preserveAspectRatio=True,
                mask="auto",
            )
        except Exception:
            pass

    canvas.setFont("Helvetica", 7)

    canvas.drawCentredString(
        width / 2,
        10 * mm,
        f"LegalEase | "
        f"{brand_name or 'AI-assisted draft'} | "
        f"Page {document.page}",
    )

    canvas.restoreState()


def make_pdf(
    text: str,
    doc_type: str,
    brand_name: Optional[str] = None,
    logo_path: Optional[str] = None,
) -> bytes:
    buffer = BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=32 * mm,
        bottomMargin=18 * mm,
        title=doc_type,
        author="LegalEase",
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "LegalTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        alignment=TA_CENTER,
        spaceAfter=10,
    )

    heading_style = ParagraphStyle(
        "LegalHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        spaceBefore=8,
        spaceAfter=5,
    )

    body_style = ParagraphStyle(
        "LegalBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10.5,
        leading=15,
        alignment=TA_LEFT,
        spaceAfter=7,
    )

    story = []

    story.append(
        Paragraph(
            escape(doc_type.upper()),
            title_style,
        )
    )

    if brand_name:
        brand_style = ParagraphStyle(
            "BrandStyle",
            parent=body_style,
            alignment=TA_CENTER,
            fontName="Helvetica-Bold",
        )

        story.append(
            Paragraph(
                escape(brand_name),
                brand_style,
            )
        )

        story.append(Spacer(1, 6))

    for line in split_document_lines(text):
        escaped = escape(line)

        style = (
            heading_style
            if is_heading(line)
            else body_style
        )

        story.append(
            Paragraph(
                escaped,
                style,
            )
        )

    document.build(
        story,
        onFirstPage=lambda canvas, doc: pdf_header_footer(
            canvas,
            doc,
            brand_name,
            logo_path,
        ),
        onLaterPages=lambda canvas, doc: pdf_header_footer(
            canvas,
            doc,
            brand_name,
            logo_path,
        ),
    )

    return buffer.getvalue()